"""Excel export implementation using openpyxl."""

import io

from openpyxl import Workbook

from app.platform.exports.csv_exporter import EXPORT_COLUMNS


def export_excel(companies: list[dict]) -> bytes:
    """Export companies to Excel workbook bytes with multiple sheets."""
    wb = Workbook()

    # Master Universe sheet
    ws_master = wb.active
    ws_master.title = "Master Universe"
    ws_master.append(EXPORT_COLUMNS)
    for company in companies:
        row = []
        for col in EXPORT_COLUMNS:
            val = company.get(col, "")
            if isinstance(val, list):
                val = "; ".join(str(v) for v in val)
            row.append(val)
        ws_master.append(row)

    # Priority Outreach sheet
    ws_outreach = wb.create_sheet("Priority Outreach")
    outreach_cols = ["canonical_name", "hq_state", "ownership_tier", "revenue_estimate", "total_score", "website"]
    ws_outreach.append(outreach_cols)
    outreach = [c for c in companies if c.get("eligible_for_outreach")]
    outreach.sort(key=lambda c: c.get("total_score", 0), reverse=True)
    for company in outreach:
        row = [company.get(col, "") for col in outreach_cols]
        ws_outreach.append(row)

    # Capital Structure sheet
    ws_capstruct = wb.create_sheet("Capital Structure")
    cap_cols = [
        "canonical_name", "pb_entity_id", "facility_type", "facility_amount",
        "pricing", "lender_names", "close_date", "maturity_date",
    ]
    ws_capstruct.append(cap_cols)
    for company in companies:
        if company.get("has_debt"):
            row = []
            for col in cap_cols:
                val = company.get(col, "")
                if isinstance(val, list):
                    val = "; ".join(str(v) for v in val)
                row.append(val)
            ws_capstruct.append(row)

    # Enrichment Sources sheet
    ws_enrichment = wb.create_sheet("Enrichment Sources")
    enrichment_cols = [
        "canonical_name", "bizapi_status", "bizapi_duns", "bizapi_match_method",
        "bizapi_match_confidence", "bizapi_verified_name", "naics_code", "naics_description",
        "sic_code", "bizapi_sales_volume", "bizapi_employee_count",
        "ciq_status", "ciq_entity_id", "ciq_revenue", "ciq_ebitda",
        "ciq_total_debt", "ciq_ownership_type",
        "pb_status", "pb_entity_id",
        "revenue_source", "ownership_source",
    ]
    ws_enrichment.append(enrichment_cols)
    for company in companies:
        row = []
        for col in enrichment_cols:
            val = company.get(col, "")
            if isinstance(val, list):
                val = "; ".join(str(v) for v in val)
            row.append(val)
        ws_enrichment.append(row)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
